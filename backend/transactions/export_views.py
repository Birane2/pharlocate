"""
Export views for transactions: PDF, Excel, Word.
Pharmacien: only their pharmacy's transactions.
Admin: all transactions (with current filters applied).
"""
import io
from datetime import datetime

from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status as http_status

from config.permissions import IsAdminRole, IsAuthenticatedWithTokenMessage, IsPharmacien

from .models import Transaction
from .views import build_admin_financial_transactions, _full_name

# ── Shared constants ──────────────────────────────────────────────────────────

_TYPE_LABELS = {
    'paiement': 'Paiement réservation',
    'commission': 'Commission',
    'remboursement': 'Remboursement',
    'ajustement': 'Ajustement',
    'abonnement': 'Abonnement',
}

PHARMACIEN_HEADERS = [
    'Référence', 'Type', 'Client', 'Méthode',
    'Montant brut', 'Commission', 'Net pharmacie', 'Date',
]

ADMIN_HEADERS = [
    'Référence', 'Type', 'Pharmacie', 'Client',
    'Méthode', 'Total', 'Commission', 'Statut', 'Date',
]

BRAND_BLUE = '#2F6E9E'
BRAND_BG = '#F8FAFC'


def _fmt_money(value):
    try:
        return f'{float(value):,.0f} MRU'
    except (TypeError, ValueError):
        return '0 MRU'


def _fmt_date(dt):
    if not dt:
        return ''
    return dt.strftime('%d/%m/%Y %H:%M')


def _generation_label():
    return f'Généré le {datetime.now().strftime("%d/%m/%Y à %H:%M")}'


# ── Data builders ─────────────────────────────────────────────────────────────

def _get_pharmacien_rows(request):
    pharmacy = getattr(request.user, 'pharmacy', None)
    if not pharmacy:
        return []

    params = request.query_params
    qs = (
        Transaction.objects
        .filter(pharmacy=pharmacy)
        .select_related('payment__payment_method', 'user')
        .order_by('-date_creation')
    )

    if params.get('start_date'):
        qs = qs.filter(date_creation__date__gte=params['start_date'])
    if params.get('end_date'):
        qs = qs.filter(date_creation__date__lte=params['end_date'])
    if params.get('type_transaction'):
        qs = qs.filter(type_transaction=params['type_transaction'])
    if params.get('search'):
        qs = qs.filter(reference_transaction__icontains=params['search'])

    rows = []
    for t in qs:
        rows.append([
            t.reference_transaction or '',
            _TYPE_LABELS.get(t.type_transaction, t.type_transaction),
            _full_name(t.user),
            (t.payment.payment_method.nom
             if t.payment_id and t.payment.payment_method_id
             else ''),
            _fmt_money(t.montant_brut),
            _fmt_money(t.commission),
            _fmt_money(t.montant_pharmacie),
            _fmt_date(t.date_creation),
        ])
    return rows


def _get_admin_rows(request):
    items = build_admin_financial_transactions(request)
    rows = []
    for item in items:
        rows.append([
            item.get('reference', ''),
            item.get('type_label', ''),
            item.get('pharmacy_name', ''),
            item.get('user_name', ''),
            item.get('payment_method', ''),
            _fmt_money(item.get('total_amount', 0)),
            _fmt_money(item.get('platform_commission', 0)),
            item.get('status_label', ''),
            _fmt_date(item.get('created_at')),
        ])
    return rows


# ── PDF builder ───────────────────────────────────────────────────────────────

def _build_pdf(doc_title, headers, rows):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
    )
    from reportlab.lib.units import cm

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    brand_color = colors.HexColor(BRAND_BLUE)
    bg_color = colors.HexColor(BRAND_BG)
    alt_row = colors.HexColor('#EBF3F9')

    brand_style = ParagraphStyle(
        'Brand',
        parent=styles['Title'],
        textColor=brand_color,
        fontSize=16,
        spaceAfter=2,
    )
    sub_style = ParagraphStyle(
        'Sub',
        parent=styles['Heading2'],
        fontSize=12,
        spaceAfter=2,
    )
    meta_style = ParagraphStyle(
        'Meta',
        parent=styles['Normal'],
        fontSize=8,
        textColor=colors.HexColor('#6B7280'),
        spaceAfter=8,
    )

    story = [
        Paragraph('PharmaLocate', brand_style),
        Paragraph(doc_title, sub_style),
        Paragraph(_generation_label(), meta_style),
        Spacer(1, 0.4 * cm),
    ]

    data = [headers] + (rows or [['Aucune donnée disponible.']])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        # Header
        ('BACKGROUND', (0, 0), (-1, 0), brand_color),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        # Body
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, alt_row]),
        ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))

    story.append(table)
    doc.build(story)
    buffer.seek(0)
    return buffer


# ── Excel builder ─────────────────────────────────────────────────────────────

def _build_excel(doc_title, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Transactions'

    col_count = len(headers)
    last_col = get_column_letter(col_count)

    # Title merged row
    ws.merge_cells(f'A1:{last_col}1')
    ws['A1'] = f'PharmaLocate — {doc_title}'
    ws['A1'].font = Font(bold=True, size=14, color='2F6E9E')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 22

    # Generated label
    ws.merge_cells(f'A2:{last_col}2')
    ws['A2'] = _generation_label()
    ws['A2'].font = Font(italic=True, size=9, color='6B7280')
    ws['A2'].alignment = Alignment(horizontal='center')

    thin_side = Side(style='thin', color='CBD5E1')
    border = Border(bottom=thin_side, right=thin_side, left=thin_side, top=thin_side)
    header_fill = PatternFill(fgColor='2F6E9E', fill_type='solid')
    alt_fill = PatternFill(fgColor='EBF3F9', fill_type='solid')

    # Headers on row 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    ws.row_dimensions[4].height = 20

    # Data rows starting at row 5
    for row_idx, row_data in enumerate(rows or [], 5):
        fill = alt_fill if (row_idx % 2 == 0) else None
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value is not None else '')
            cell.alignment = Alignment(vertical='center')
            cell.border = border
            if fill:
                cell.fill = fill

    # Auto-width columns (iterate by index to avoid merged-cell objects)
    for col_idx in range(1, col_count + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(4, ws.max_row + 1)),
            default=8,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 3, 35)

    # Freeze header
    ws.freeze_panes = 'A5'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ── Word builder ──────────────────────────────────────────────────────────────

def _build_word(doc_title, headers, rows):
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def _set_cell_bg(cell, hex_color):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), hex_color)
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:val'), 'clear')
        tc_pr.append(shd)

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)

    # Brand heading
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('PharmaLocate')
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x2F, 0x6E, 0x9E)

    # Report title
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(doc_title)
    run2.bold = True
    run2.font.size = Pt(13)

    # Generated date
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run(_generation_label())
    run3.italic = True
    run3.font.size = Pt(9)
    run3.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph()

    if not rows:
        doc.add_paragraph('Aucune donnée disponible.')
    else:
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header row
        hdr = table.rows[0]
        for i, hdr_text in enumerate(headers):
            cell = hdr.cells[i]
            cell.text = hdr_text
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.runs[0]
            run.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _set_cell_bg(cell, '2F6E9E')

        # Data rows
        for row_idx, row_data in enumerate(rows):
            row = table.rows[row_idx + 1]
            bg = 'EBF3F9' if row_idx % 2 == 0 else 'FFFFFF'
            for col_idx, val in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.text = str(val) if val is not None else ''
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].font.size = Pt(7)
                _set_cell_bg(cell, bg)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ── Response helpers ──────────────────────────────────────────────────────────

def _pdf_response(buffer, filename):
    resp = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


def _excel_response(buffer, filename):
    resp = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


def _word_response(buffer, filename):
    resp = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


# ── Pharmacien export views ───────────────────────────────────────────────────

class PharmacienTransactionsPDFExport(APIView):
    permission_classes = [IsAuthenticatedWithTokenMessage, IsPharmacien]

    def get(self, request):
        rows = _get_pharmacien_rows(request)
        buf = _build_pdf('Transactions — Ma pharmacie', PHARMACIEN_HEADERS, rows)
        return _pdf_response(buf, 'transactions-pharmacie.pdf')


class PharmacienTransactionsExcelExport(APIView):
    permission_classes = [IsAuthenticatedWithTokenMessage, IsPharmacien]

    def get(self, request):
        rows = _get_pharmacien_rows(request)
        buf = _build_excel('Transactions — Ma pharmacie', PHARMACIEN_HEADERS, rows)
        return _excel_response(buf, 'transactions-pharmacie.xlsx')


class PharmacienTransactionsWordExport(APIView):
    permission_classes = [IsAuthenticatedWithTokenMessage, IsPharmacien]

    def get(self, request):
        rows = _get_pharmacien_rows(request)
        buf = _build_word('Transactions — Ma pharmacie', PHARMACIEN_HEADERS, rows)
        return _word_response(buf, 'transactions-pharmacie.docx')


# ── Admin export views ────────────────────────────────────────────────────────

class AdminTransactionsPDFExport(APIView):
    permission_classes = [IsAuthenticatedWithTokenMessage, IsAdminRole]

    def get(self, request):
        rows = _get_admin_rows(request)
        buf = _build_pdf('Rapport Transactions — PharmaLocate', ADMIN_HEADERS, rows)
        return _pdf_response(buf, 'transactions-admin.pdf')


class AdminTransactionsExcelExport(APIView):
    permission_classes = [IsAuthenticatedWithTokenMessage, IsAdminRole]

    def get(self, request):
        rows = _get_admin_rows(request)
        buf = _build_excel('Rapport Transactions — PharmaLocate', ADMIN_HEADERS, rows)
        return _excel_response(buf, 'transactions-admin.xlsx')


class AdminTransactionsWordExport(APIView):
    permission_classes = [IsAuthenticatedWithTokenMessage, IsAdminRole]

    def get(self, request):
        rows = _get_admin_rows(request)
        buf = _build_word('Rapport Transactions — PharmaLocate', ADMIN_HEADERS, rows)
        return _word_response(buf, 'transactions-admin.docx')
